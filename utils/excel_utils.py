from openpyxl import load_workbook, Workbook

def write_sheet(data, filepath, sheet_name):
    try:
        workbook = load_workbook(filepath)
    except FileNotFoundError:
        workbook = Workbook()

    workbook.create_sheet(sheet_name)
    worksheet = workbook[sheet_name]
    for row in data:
        worksheet.append(row)
    workbook.save(filepath)

